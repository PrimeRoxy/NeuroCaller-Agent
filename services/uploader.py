import logging
from typing import Optional
from fastapi import UploadFile
from google import genai
from google.genai import types
import os
from dotenv import load_dotenv
import aiohttp
import tempfile
import uuid
import boto3
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from io import BytesIO

logging.basicConfig(level=logging.INFO)

load_dotenv()
GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
client = genai.Client(api_key=GOOGLE_API_KEY)

S3_BUCKET_NAME = os.getenv("AWS_S3_BUCKET")
AWS_REGION = os.getenv("S3_REGION")
print(f"S3_BUCKET_NAME: {S3_BUCKET_NAME}, AWS_REGION: {AWS_REGION}")
s3_client = boto3.client("s3", region_name=AWS_REGION)

def upload_file_for_gemini(file, mime_type: str) -> types.File:
    try:
        uploaded = client.files.upload(
            file=file, config=types.UploadFileConfig(mime_type=mime_type)
        )
        return uploaded
    except Exception as e:
        logging.error(f"Gemini upload error: {e}")
        return None

async def upload_url_to_s3(recording_url: str, headers: Optional[dict] = None) -> str:
    try:
        logging.info(f"Downloading recording from URL: {recording_url}")
        connector = aiohttp.TCPConnector(ssl=False)
        async with aiohttp.ClientSession(connector=connector) as session:
            async with session.get(recording_url, headers=headers) as response:
                if response.status != 200:
                    error_msg = f"Failed to download recording: HTTP {response.status}"
                    logging.error(error_msg)
                    return "Failed to download recording"
                file_content = await response.read()

        if not file_content:
            logging.error("Downloaded file is empty.")
            return "Downloaded file is empty"

        if not S3_BUCKET_NAME or not AWS_REGION:
            logging.error("Missing S3 configuration.")
            return "S3 configuration error"
        file_ext = ".mp3"
        file_id = f"{uuid.uuid4()}{file_ext}"
        
        s3_client.put_object(
            Bucket=S3_BUCKET_NAME,
            Key=file_id,
            Body=file_content,
            ContentType="audio/mpeg"
        )
        
        s3_url = f"https://{S3_BUCKET_NAME}.s3.{AWS_REGION}.amazonaws.com/{file_id}"
        return s3_url

    except Exception as e:
        error_msg = f"Upload failed on S3 bucket : {e}"
        logging.error(error_msg)
        return "Upload failed on S3 bucket"

def convert_docx_to_pdf(file_obj):
    """
    Converts a DOCX file object to a temporary PDF file using ReportLab.
    Returns a file object ready to upload.
    """
    try:
        file_obj.seek(0)
        doc = Document(file_obj)

        styles = getSampleStyleSheet()
        story = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                story.append(Paragraph(text, styles['Normal']))
                story.append(Spacer(1, 10))

        for table in doc.tables:
            data = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if data:
                table_obj = Table(data)
                table_obj.setStyle([
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP')
                ])
                story.append(table_obj)
                story.append(Spacer(1, 10))

        tmp_path = tempfile.mktemp(suffix=".pdf")
        doc_pdf = SimpleDocTemplate(tmp_path, pagesize=A4)
        doc_pdf.build(story)

        return open(tmp_path, "rb")

    except Exception as e:
        logging.error(f"[DOCX→PDF] Conversion error: {e}")
        return None
